from __future__ import annotations

import math
from pathlib import Path
import re
from typing import Iterable

from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from marvis.artifacts import ArtifactUnitOfWork
from marvis.formatting import period_text as _period_text
from marvis.formatting import psi_reference_month_text as _psi_reference_month_text
from marvis.formatting import ratio as _ratio
from marvis.formatting import score_interval as _score_interval
from marvis.output.image_render import render_roc_ks_graph
from marvis.output.styles import (
    BORDER_COLOR,
    BRAND_HEADER_FILL,
    BRAND_HEADER_FONT_COLOR,
    FONT_NAME,
    FONT_SIZE_PT,
    STRESS_MEDIUM_FILL,
    ks_drop_ratio,
    stress_ks_risk,
    stress_psi_risk,
    stress_risk_cell_color,
    status_cell_color,
    worst_stress_risk,
)
from marvis.output.xlsx_safety import safe_xlsx_cell
from marvis.validation.results import (
    BinRow,
    ConsistencyStatus,
    ValidationResults,
)

_INVALID_SHEET_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")
_MAX_SHEET_TITLE_LENGTH = 31


def write_validation_excel(
    results: ValidationResults,
    output_path: Path,
    *,
    report_values: dict[str, object] | None = None,
    image_output_dir: Path | None = None,
) -> Path:
    """Write the validation workbook and its rendered chart images.

    Normally this function owns an atomic workbook/image transaction.  A
    caller that already owns a larger artifact transaction may instead supply
    ``image_output_dir`` alongside a staged ``output_path``.  In that mode the
    caller owns promotion and rollback of both paths; creating a nested unit
    of work would otherwise leave chart images behind when the outer database
    transaction fails.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    owns_artifacts = image_output_dir is None
    uow: ArtifactUnitOfWork | None = None
    if owns_artifacts:
        uow = ArtifactUnitOfWork()
        workbook_artifact = uow.stage_file(output_path.parent, output_path.name)
        image_artifact = uow.stage_directory(output_path.parent, "excel_images")
        workbook_output_path = workbook_artifact.path
        chart_image_dir = image_artifact.path
    else:
        workbook_artifact = None
        workbook_output_path = output_path
        chart_image_dir = Path(image_output_dir)
        chart_image_dir.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    try:
        _write_overview(workbook, results)
        _write_report_texts(workbook, report_values)
        _write_basic_info(workbook, results)
        _write_monthly_distribution(workbook, results)
        _write_hyperparameters(workbook, results)
        _write_feature_importance(workbook, results)
        _write_effectiveness_overall(workbook, results)
        _write_psi_stability(workbook, results)
        _write_roc_ks_images(workbook, results, chart_image_dir)
        for split in ("train", "test", "oot"):
            _write_bins(
                workbook,
                f"Box_{split}",
                results.effectiveness.bin_tables.get(split, []),
                first_header=f"{split}(According totrainBox)",
            )
        independent = results.effectiveness.independent_quantile_bin_tables or {}
        if any(independent.get(split) for split in ("train", "test", "oot")):
            for split in ("train", "test", "oot"):
                _write_bins(
                    workbook,
                    f"Box_Independence 10 points_{split}",
                    independent.get(split, []),
                    first_header=f"{split}(Independence 10 points)",
                )
        _write_monthly_effectiveness(workbook, results)
        _write_stress_summary(workbook, results)
        for category_result in results.stress_test.per_category:
            sheet_name = f"Pressure test_Box_{category_result.category}"
            _write_bins(
                workbook,
                sheet_name,
                category_result.bin_table,
                first_header=category_result.category,
                include_bin_share=True,
            )

        workbook.save(workbook_output_path)
        if uow is None:
            return workbook_output_path
        uow.promote_all()
        uow.commit()
        return workbook_artifact.final_path
    except Exception:
        if uow is not None:
            uow.rollback()
        raise


def _write_overview(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Authentication Overview")
    rows: list[tuple[object, object]] = [
        ("Item", "Value"),
        ("Model", f"{results.model_name} {results.model_version}"),
        ("Algorithm", results.algorithm),
    ]
    if results.pmml_scoring is not None:
        scoring = results.pmml_scoring
        rows.extend(
            [
                ("PMMLRating status", scoring.status),
                ("Full sample", scoring.input_row_count),
                ("Success", scoring.success_count),
                ("Failed scoring", scoring.failure_count),
                ("Empty output", scoring.null_count),
                ("Non-limited output", scoring.non_finite_count),
                ("PMMLOutput Field", scoring.output_field),
                ("Batch Engine", scoring.engine),
                ("Time-shaping.(sec)", scoring.elapsed_seconds),
                ("Swallow.(Okay./sec)", scoring.rows_per_second),
            ]
        )
        status_value = (
            ConsistencyStatus.PASS
            if scoring.status == "pass"
            else ConsistencyStatus.FAIL
            if scoring.status == "fail"
            else ConsistencyStatus.REVIEW
        )
    elif results.reproducibility is not None:
        summary = results.reproducibility.summary
        rows.extend(
            [
                ("Sample rows", results.reproducibility.sample_size),
                ("Lines", summary.match_count),
                ("Number of lines of variance", summary.mismatch_count),
                ("Max. Absolute", summary.max_abs_diff),
                ("Recoverability status", summary.status.value),
            ]
        )
        status_value = summary.status
    else:
        rows.append(("Authentication status", "Lack of scoring or duplicate evidence"))
        status_value = ConsistencyStatus.REVIEW
    _write_rows(sheet, rows, header_rows=1)
    sheet.cell(row=len(rows), column=2).fill = PatternFill(
        start_color=status_cell_color(status_value),
        end_color=status_cell_color(status_value),
        fill_type="solid",
    )


def _write_report_texts(
    workbook: Workbook,
    report_values: dict[str, object] | None,
) -> None:
    if not report_values:
        return
    rows: list[tuple[object, object]] = [("Report Fields", "Contents")]
    rows.extend(
        (str(key), value)
        for key, value in sorted(report_values.items())
        if value not in (None, "")
    )
    if len(rows) == 1:
        return
    sheet = workbook.create_sheet("Text of the report")
    _write_rows(sheet, rows, header_rows=1)
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(vertical="top", wrap_text=True)


def _write_basic_info(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Sample Basic Information")
    total_count = sum(row.sample_count for row in results.basic_info.split_summary)
    rows: list[tuple] = [("Dataset", "Time frame", "Sample Volume", "Samples as a percentage", "Bad sample quantity", "Overdue rate")]
    rows.extend(
        (
            row.split,
            _period_text(row.period_start, row.period_end, default="-"),
            row.sample_count,
            _ratio(row.sample_count, total_count),
            row.bad_count,
            row.bad_rate,
        )
        for row in results.basic_info.split_summary
    )
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        percent_columns={3, 5},
        data_bar_columns={2: "5A8AC6", 5: "F8696B"},
    )


def _write_monthly_distribution(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Samples distributed monthly")
    total_count = sum(row.sample_count for row in results.basic_info.monthly_distribution)
    rows: list[tuple] = [("Month", "Sample Volume", "Samples as a percentage", "Bad sample quantity", "Overdue rate")]
    rows.extend(
        (row.month, row.sample_count, _ratio(row.sample_count, total_count), row.bad_count, row.bad_rate)
        for row in results.basic_info.monthly_distribution
    )
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        percent_columns={2, 4},
        data_bar_columns={1: "5A8AC6", 4: "F8696B"},
    )


def _write_hyperparameters(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Model superb")
    rows: list[tuple] = [("Parameters", "Value")]
    rows.extend((str(key), str(value))
                for key, value in results.basic_info.hyperparameters.items())
    _write_rows(sheet, rows, header_rows=1)


def _write_feature_importance(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Characteristic importance")
    rows: list[tuple] = [("Rank", "Characteristics", "Category", "Importance")]
    rows.extend((row.rank, row.feature, row.category, row.importance)
                for row in results.basic_info.feature_importance)
    _write_rows(sheet, rows, header_rows=1, decimal_columns={3})


def _write_effectiveness_overall(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Model Effects")
    split_summary = {row.split: row for row in results.basic_info.split_summary}
    rows: list[tuple] = [(
        "Dataset", "Time frame", "Sample Volume", "Overdue rate", "Bad sample quantity",
        "KS(%)", "AUC(%)", "5%Headlift", "5%Endlift", "PSI",
    )]
    rows.extend(
        (
            row.split,
            _period_text(
                split_summary.get(row.split).period_start if split_summary.get(row.split) else "",
                split_summary.get(row.split).period_end if split_summary.get(row.split) else "",
                default="-",
            ),
            row.sample_count,
            row.bad_rate,
            _bad_count(row),
            _pct_point(row.ks),
            _pct_point(row.auc),
            _optional_number(row.head_lift_5pct),
            _optional_number(row.tail_lift_5pct),
            "BASE" if row.split == "train" else row.psi_vs_train,
        )
        for row in results.effectiveness.overall
    )
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        percent_columns={3},
        number_formats={5: "0.0", 6: "0.0", 7: "0.00", 8: "0.00", 9: "0.000"},
        data_bar_columns={5: "63BE7B"},
    )


def _write_psi_stability(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("PSIStability")
    rows: list[tuple] = [(
        "Box", "train+testNumber of samples", "train+testPercentage", "ootNumber of samples", "ootPercentage", "PSI",
    )]
    rows.extend(
        (
            row.bin_label,
            row.expected_count,
            row.expected_pct,
            row.actual_count,
            row.actual_pct,
            row.psi,
        )
        for row in results.effectiveness.psi_stability_table
    )
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        percent_columns={2, 4},
        number_formats={5: "0.0000"},
        data_bar_columns={1: "5A8AC6", 3: "5A8AC6"},
        color_scale_columns={5},
    )


def _write_roc_ks_images(workbook: Workbook, results: ValidationResults, image_dir: Path) -> None:
    sheet = workbook.create_sheet("ROC_KSCurve")
    image_dir.mkdir(parents=True, exist_ok=True)
    for index, split in enumerate(("train", "test", "oot")):
        start_row = index * 32 + 1
        sheet.cell(
            row=start_row,
            column=1,
            value=safe_xlsx_cell(f"{split} ROCCurves andKSCurve"),
        )
        sheet.cell(row=start_row, column=1).font = Font(name=FONT_NAME, size=FONT_SIZE_PT, bold=True)
        image_path = render_roc_ks_graph(
            results.effectiveness.roc_ks_curves.get(split),
            image_dir / f"roc_ks_graph_{split}.png",
            title_prefix=split,
        )
        image = OpenpyxlImage(str(image_path))
        image.width = 720
        image.height = 480
        sheet.add_image(image, f"A{start_row + 1}")
    sheet.column_dimensions["A"].width = 100


def _write_monthly_effectiveness(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Monthly effect")
    by_month: dict[str, dict[str, object]] = {}
    for ks_row in results.effectiveness.monthly_ks:
        by_month.setdefault(ks_row.month, {})["KS"] = ks_row.ks
        by_month[ks_row.month]["Number of samples"] = ks_row.sample_count
        by_month[ks_row.month]["Bad sample quantity"] = ks_row.bad_count
        by_month[ks_row.month]["Overdue rate"] = ks_row.bad_rate
        by_month[ks_row.month]["AUC"] = ks_row.auc
        by_month[ks_row.month]["5%Headlift"] = ks_row.head_lift_5pct
        by_month[ks_row.month]["5%Endlift"] = ks_row.tail_lift_5pct
    for psi_row in results.effectiveness.monthly_psi:
        by_month.setdefault(psi_row.month, {})["PSI"] = psi_row.psi_vs_train
        by_month[psi_row.month]["PSI(First month benchmark)"] = psi_row.psi_first_month
        by_month[psi_row.month]["PSI(End-month benchmark)"] = psi_row.psi_last_month
        by_month[psi_row.month]["PSI(It's a lot more than the previous sample month.)"] = psi_row.psi_mom
        by_month[psi_row.month]["PSIMonth of reference"] = _psi_reference_month_text(
            psi_row.psi_mom_reference_month,
            has_calendar_gap=psi_row.psi_mom_has_calendar_gap,
        )
    months = sorted(by_month)
    first_month = months[0] if months else ""
    last_month = months[-1] if months else ""
    rows: list[tuple] = [(
        "Month", "Sample Volume", "Overdue rate", "Bad sample quantity", "KS(%)", "AUC(%)",
        "5%Headlift", "5%Endlift", "PSI(First month benchmark)", "PSI(End-month benchmark)",
        "PSI(It's a lot more than the previous sample month.)", "PSIMonth of reference",
    )]
    rows.extend(
        (
            month,
            data.get("Number of samples", 0),
            data.get("Overdue rate", 0.0),
            data.get("Bad sample quantity", 0),
            _pct_point(float(data.get("KS", 0.0))),
            _pct_point(float(data.get("AUC", 0.0))),
            _optional_number(data.get("5%Headlift")),
            _optional_number(data.get("5%Endlift")),
            "BASE" if month == first_month else _optional_number(data.get("PSI(First month benchmark)")),
            "BASE" if month == last_month else _optional_number(data.get("PSI(End-month benchmark)")),
            "-" if month == first_month else _optional_number(data.get("PSI(It's a lot more than the previous sample month.)")),
            "-" if month == first_month else data.get("PSIMonth of reference", ""),
        )
        for month, data in sorted(by_month.items())
    )
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        percent_columns={2},
        number_formats={4: "0.0", 5: "0.0", 6: "0.00", 7: "0.00", 8: "0.000", 9: "0.000", 10: "0.000"},
        data_bar_columns={4: "63BE7B"},
    )


def _write_stress_summary(workbook: Workbook, results: ValidationResults) -> None:
    sheet = workbook.create_sheet("Pressure test_Summary")
    sheet.sheet_view.showGridLines = False
    rows: list[tuple] = [(
        "Category",
        "KS_baseline",
        "KS_after",
        "KS_delta",
        "KSDeclination rate",
        "PSI",
        "Test Results",
    )]
    baseline_ks = results.stress_test.baseline.ks
    baseline_psi = _oot_psi(results)
    rows.append(("baseline", baseline_ks, "", "", "", baseline_psi, ""))
    category_risks: list[str | None] = []
    for item in results.stress_test.per_category:
        ks_risk = (
            stress_ks_risk(baseline_ks, item.ks_after)
            if item.ks_after is not None and item.status == "completed"
            else None
        )
        psi_risk = (
            stress_psi_risk(item.psi_vs_baseline)
            if item.status == "completed"
            else None
        )
        overall_risk = worst_stress_risk(ks_risk, psi_risk)
        category_risks.append(overall_risk)
        rows.append((
            item.category,
            baseline_ks,
            item.ks_after if item.ks_after is not None else "",
            item.ks_delta if item.ks_delta is not None else "",
            (
                ks_drop_ratio(baseline_ks, item.ks_after)
                if item.ks_after is not None
                else ""
            ),
            item.psi_vs_baseline if item.psi_vs_baseline is not None else "",
            _stress_risk_label(overall_risk),
        ))

    _write_rows(
        sheet,
        rows,
        header_rows=1,
        decimal_columns={1, 2, 3, 5},
        number_formats={4: "0.0%"},
    )
    sheet.freeze_panes = "A2"

    if baseline_psi is not None:
        _fill_stress_risk_cell(sheet.cell(row=2, column=6), stress_psi_risk(baseline_psi))
    for row_index, (item, overall_risk) in enumerate(
        zip(results.stress_test.per_category, category_risks, strict=True),
        start=3,
    ):
        ks_risk = (
            stress_ks_risk(baseline_ks, item.ks_after)
            if item.ks_after is not None and item.status == "completed"
            else None
        )
        psi_risk = (
            stress_psi_risk(item.psi_vs_baseline)
            if item.status == "completed"
            else None
        )
        _fill_stress_risk_cell(sheet.cell(row=row_index, column=4), ks_risk)
        _fill_stress_risk_cell(sheet.cell(row=row_index, column=5), ks_risk)
        _fill_stress_risk_cell(sheet.cell(row=row_index, column=6), psi_risk)
        if overall_risk is None:
            sheet.cell(row=row_index, column=7).fill = PatternFill(
                start_color=STRESS_MEDIUM_FILL,
                end_color=STRESS_MEDIUM_FILL,
                fill_type="solid",
            )
        else:
            _fill_stress_risk_cell(sheet.cell(row=row_index, column=7), overall_risk)

    notes_start = len(rows) + 2
    unclassified = results.stress_test.unclassified_features
    coverage_text = f"Unclassified features{len(unclassified)} individual"
    if unclassified:
        coverage_text += ":" + _feature_name_preview(unclassified)
    source_counts = ",".join(
        f"{source} {count}"
        for source, count in sorted(results.stress_test.category_source_counts.items())
    )
    if source_counts:
        coverage_text = f"{coverage_text};Category map source:{source_counts}"
    note_rows = [
        (
            "Test Method",
            "Pressure tests are primarily based on the deviation of key indicators of the model under a missing source scenario and their effects,"
            "Focus on effectivenessKS,StabilityPSI andOOT .",
        ),
        (
            f"Class coverage:{_stress_status_label(results.stress_test.status)}",
            coverage_text,
        ),
        (
            "Set-9999 Annotations",
            "Original fields of the module are unified for each category-9999,And the same.OOT • Rescoring on the sample;"
            "Reuse All Categorybaseline OOT the same frequency-box boundary.",
        ),
    ]
    _write_note_rows(sheet, note_rows, start_row=notes_start)

    matrix_start = notes_start + len(note_rows) + 1
    alignment_notes = _write_stress_bin_share_matrix(
        sheet,
        results,
        category_risks,
        start_row=matrix_start,
    )
    legend_start = matrix_start + len(results.stress_test.baseline.bin_table) + 4
    if alignment_notes:
        for offset, note in enumerate(alignment_notes):
            sheet.cell(
                row=legend_start + offset,
                column=1,
                value=safe_xlsx_cell(note),
            )
        legend_start += len(alignment_notes) + 1
    _write_stress_legend(sheet, start_row=legend_start)

    sheet.column_dimensions["A"].width = 20
    for column in range(2, max(8, len(results.stress_test.per_category) + 2)):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{get_column_letter(max(7, len(results.stress_test.per_category) + 1))}{sheet.max_row}"


def _oot_psi(results: ValidationResults) -> float | None:
    for row in results.effectiveness.overall:
        if str(row.split).lower() == "oot":
            return float(row.psi_vs_train)
    return None


def _stress_risk_label(risk: str | None) -> str:
    return {
        "low": "Low risk",
        "medium": "Medium risk",
        "high": "High risk",
    }.get(str(risk or ""), "Unable to assess")


def _fill_stress_risk_cell(cell, risk: str | None) -> None:
    color = stress_risk_cell_color(risk)
    if not color:
        return
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")


def _write_note_rows(sheet, rows: list[tuple[str, str]], *, start_row: int) -> None:
    label_font = Font(name=FONT_NAME, size=FONT_SIZE_PT, bold=True, color="C00000")
    body_font = Font(name=FONT_NAME, size=FONT_SIZE_PT)
    for offset, (label, body) in enumerate(rows):
        row_index = start_row + offset
        label_cell = sheet.cell(
            row=row_index,
            column=1,
            value=safe_xlsx_cell(label),
        )
        body_cell = sheet.cell(
            row=row_index,
            column=2,
            value=safe_xlsx_cell(body),
        )
        label_cell.font = label_font
        body_cell.font = body_font
        label_cell.alignment = Alignment(vertical="top", wrap_text=True)
        body_cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=7)
        sheet.row_dimensions[row_index].height = 34


def _write_stress_bin_share_matrix(
    sheet,
    results: ValidationResults,
    category_risks: list[str | None],
    *,
    start_row: int,
) -> list[str]:
    baseline_bins = results.stress_test.baseline.bin_table
    categories = results.stress_test.per_category
    aligned_shares: list[list[float] | None] = []
    notes: list[str] = []
    for item in categories:
        shares = _aligned_stress_bin_shares(baseline_bins, item.bin_table)
        if item.status != "completed" or shares is None:
            aligned_shares.append(None)
            reason = item.error or "The border of the box is unmatched."
            notes.append(f"{item.category}:{reason},The share of the boxes remains empty.")
        else:
            aligned_shares.append(shares)

    rows: list[tuple] = [("OOTBox", *(item.category for item in categories))]
    for bin_position, baseline_bin in enumerate(baseline_bins):
        rows.append((
            _score_interval(baseline_bin.score_lower, baseline_bin.score_upper),
            *(
                shares[bin_position] if shares is not None else ""
                for shares in aligned_shares
            ),
        ))
    rows.append((
        "Total",
        *(sum(shares) if shares is not None else "" for shares in aligned_shares),
    ))
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        number_formats={index: "0.0%" for index in range(1, len(categories) + 1)},
        start_row=start_row,
    )
    for category_index, risk in enumerate(category_risks, start=1):
        color = stress_risk_cell_color(risk) or "BFBFBF"
        _apply_reference_conditional_formatting(
            sheet,
            start_row=start_row + 1,
            end_row=start_row + len(baseline_bins),
            data_bar_columns={category_index: color},
            color_scale_columns=set(),
        )
    return notes


def _aligned_stress_bin_shares(
    baseline_bins: list[BinRow],
    category_bins: list[BinRow],
) -> list[float] | None:
    if len(baseline_bins) != len(category_bins):
        return None
    for baseline, category in zip(baseline_bins, category_bins, strict=True):
        if baseline.bin_index != category.bin_index:
            return None
        if not _same_score_bound(baseline.score_lower, category.score_lower):
            return None
        if not _same_score_bound(baseline.score_upper, category.score_upper):
            return None
    total = sum(row.sample_count for row in category_bins)
    if total <= 0:
        return None
    return [row.sample_count / total for row in category_bins]


def _same_score_bound(left: float, right: float) -> bool:
    return math.isclose(float(left), float(right), rel_tol=1e-9, abs_tol=1e-12)


def _write_stress_legend(sheet, *, start_row: int) -> None:
    headers = ("Pressure test", "Low risk", "Medium risk", "High risk")
    descriptions = (
        "Validity of pressure test focus when source is missingKS,StabilityPSI andOOT .",
        "✓ KSDeclining range:[0,10%)\nPSIScope:[0,0.10)\nThe absence of features has little impact and is usable, but requires routine stability monitoring.",
        "— KSDeclining range:[10%,20%)\nPSIScope:[0.10,0.25)\nEarly warning thresholds are met, which require manual verification of anomalies and timely notification to the modelling team.",
        "! KSDeclining range:[20%,+∞)\nPSIScope:[0.25,+∞)\nReal-time monitoring and melting mechanisms are required, downgraded to a back-up rating or manual review is initiated when necessary.",
    )
    header_fill = PatternFill(
        start_color=BRAND_HEADER_FILL,
        end_color=BRAND_HEADER_FILL,
        fill_type="solid",
    )
    risk_fills = (None, "low", "medium", "high")
    for column_index, header in enumerate(headers, start=1):
        header_cell = sheet.cell(row=start_row, column=column_index, value=header)
        header_cell.font = Font(
            name=FONT_NAME,
            size=FONT_SIZE_PT,
            bold=True,
            color=BRAND_HEADER_FONT_COLOR,
        )
        header_cell.fill = header_fill
        header_cell.alignment = Alignment(horizontal="center", vertical="center")
        body_cell = sheet.cell(
            row=start_row + 1,
            column=column_index,
            value=descriptions[column_index - 1],
        )
        body_cell.font = Font(name=FONT_NAME, size=FONT_SIZE_PT)
        body_cell.alignment = Alignment(vertical="top", wrap_text=True)
        _fill_stress_risk_cell(body_cell, risk_fills[column_index - 1])
        sheet.column_dimensions[get_column_letter(column_index)].width = 28
    sheet.row_dimensions[start_row + 1].height = 86


def _feature_name_preview(features: list[str], *, limit: int = 20) -> str:
    visible = features[:limit]
    text = ",".join(visible)
    if len(features) > limit:
        return f"{text} Wait.{len(features)} individual"
    return text


def _stress_status_label(status: str) -> str:
    return {
        "completed": "Completed",
        "skipped": "Skip",
        "error": "Unusual",
        "partial": "Partially completed",
        "failed": "Failed",
    }.get(str(status or ""), str(status or ""))


def _write_bins(
    workbook: Workbook,
    sheet_name: str,
    bins: list[BinRow],
    *,
    first_header: str = "Box",
    include_bin_share: bool = False,
) -> None:
    sheet = workbook.create_sheet(_safe_sheet_title(workbook, sheet_name))
    if include_bin_share:
        rows: list[tuple] = [(
            first_header, "Total sample", "Share of boxes", "Cumulative share", "Overdue",
            "Overdue rate", "Cumulative overdue rate", "Single grouplift", "Cumulativelift", "ks",
        )]
        percent_columns = {2, 3, 5, 6}
        number_formats = {7: "0.00", 8: "0.00", 9: "0.0000"}
        color_scale_columns = {5}
        data_bar_columns = {8: "63BE7B"}
    else:
        rows = [(
            first_header, "Total sample", "Cumulative share", "Overdue", "Overdue rate",
            "Cumulative overdue rate", "Single grouplift", "Cumulativelift", "ks",
        )]
        percent_columns = {2, 4, 5}
        number_formats = {6: "0.00", 7: "0.00", 8: "0.0000"}
        color_scale_columns = {4}
        data_bar_columns = {7: "63BE7B"}
    rows.extend(_reference_bin_rows(bins, include_bin_share=include_bin_share))
    _write_rows(
        sheet,
        rows,
        header_rows=1,
        percent_columns=percent_columns,
        number_formats=number_formats,
        color_scale_columns=color_scale_columns,
        data_bar_columns=data_bar_columns,
    )


def _write_rows(
    sheet,
    rows: Iterable[tuple],
    *,
    header_rows: int,
    percent_columns: set[int] | None = None,
    decimal_columns: set[int] | None = None,
    number_formats: dict[int, str] | None = None,
    data_bar_columns: dict[int, str] | None = None,
    color_scale_columns: set[int] | None = None,
    start_row: int = 1,
) -> None:
    percent_columns = percent_columns or set()
    decimal_columns = decimal_columns or set()
    number_formats = number_formats or {}
    data_bar_columns = data_bar_columns or {}
    color_scale_columns = color_scale_columns or set()
    rows = list(rows)
    if not rows:
        return

    header_fill = PatternFill(
        start_color=BRAND_HEADER_FILL,
        end_color=BRAND_HEADER_FILL,
        fill_type="solid",
    )
    header_font = Font(
        name=FONT_NAME, size=FONT_SIZE_PT, bold=True,
        color=BRAND_HEADER_FONT_COLOR,
    )
    body_font = Font(name=FONT_NAME, size=FONT_SIZE_PT)
    thin = Side(border_style="thin", color=BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    column_count = max(len(row) for row in rows)
    header_end_row = start_row + header_rows - 1
    for row_index, row_values in enumerate(rows, start=start_row):
        for column_index in range(column_count):
            value = row_values[column_index] if column_index < len(row_values) else ""
            cell = sheet.cell(
                row=row_index,
                column=column_index + 1,
                value=safe_xlsx_cell(value),
            )
            cell.font = header_font if row_index <= header_end_row else body_font
            cell.alignment = center
            cell.border = border
            if row_index <= header_end_row:
                cell.fill = header_fill
            else:
                if column_index in percent_columns and isinstance(value, (int, float)):
                    cell.number_format = "0.00%"
                elif column_index in number_formats and isinstance(value, (int, float)):
                    cell.number_format = number_formats[column_index]
                elif column_index in decimal_columns and isinstance(value, (int, float)):
                    cell.number_format = "0.0000"

    _apply_reference_conditional_formatting(
        sheet,
        start_row=start_row + header_rows,
        end_row=start_row + len(rows) - 1,
        data_bar_columns=data_bar_columns,
        color_scale_columns=color_scale_columns,
    )

    for column_index in range(column_count):
        sheet.column_dimensions[get_column_letter(column_index + 1)].width = 14


def _apply_reference_conditional_formatting(
    sheet,
    *,
    start_row: int,
    end_row: int,
    data_bar_columns: dict[int, str],
    color_scale_columns: set[int],
) -> None:
    if end_row < start_row:
        return
    for column_index in color_scale_columns:
        column = get_column_letter(column_index + 1)
        sheet.conditional_formatting.add(
            f"{column}{start_row}:{column}{end_row}",
            ColorScaleRule(
                start_type="min", start_color="63BE7B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="F8696B",
            ),
        )
    for column_index, color in data_bar_columns.items():
        column = get_column_letter(column_index + 1)
        sheet.conditional_formatting.add(
            f"{column}{start_row}:{column}{end_row}",
            DataBarRule(
                start_type="percentile", start_value=0,
                end_type="percentile", end_value=100,
                color=color,
                showValue=True,
            ),
        )


def _reference_bin_rows(
    bins: list[BinRow],
    *,
    include_bin_share: bool = False,
) -> list[tuple]:
    total = sum(row.sample_count for row in bins)
    total_bad = sum(row.bad_count for row in bins)
    overall_bad_rate = _ratio(total_bad, total)
    cumulative_count = 0
    cumulative_bad = 0
    rows: list[tuple] = []
    for row in bins:
        cumulative_count += row.sample_count
        cumulative_bad += row.bad_count
        cumulative_bad_rate = _ratio(cumulative_bad, cumulative_count)
        values = (
            _score_interval(row.score_lower, row.score_upper),
            row.sample_count,
            _ratio(cumulative_count, total),
            row.bad_count,
            row.bad_rate,
            cumulative_bad_rate,
            row.lift,
            _ratio(cumulative_bad_rate, overall_bad_rate),
            row.ks,
        )
        if include_bin_share:
            values = values[:2] + (_ratio(row.sample_count, total),) + values[2:]
        rows.append(values)
    return rows


def _pct_point(value: float) -> float:
    return round(float(value) * 100, 1)


def _bad_count(row) -> int:
    if row.bad_count is not None:
        return int(row.bad_count)
    return int(round(row.sample_count * row.bad_rate))


def _optional_number(value) -> float | str:
    return "" if value is None else float(value)


def _safe_sheet_title(workbook: Workbook, title: str) -> str:
    base = _INVALID_SHEET_TITLE_CHARS.sub("_", str(title)).strip().strip("'") or "Sheet"
    candidate = base[:_MAX_SHEET_TITLE_LENGTH]
    if candidate not in workbook.sheetnames:
        return candidate
    index = 2
    while True:
        suffix = f"_{index}"
        candidate = f"{base[:_MAX_SHEET_TITLE_LENGTH - len(suffix)]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
        index += 1
